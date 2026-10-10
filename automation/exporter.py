"""CHANGE [2026-09-10 09:52 +08:00] [WH400]: 为明细表导出增加菜单内容校验和局部重试，阻止筛选器菜单进入下载阶段。"""

from __future__ import annotations

import argparse
import asyncio
import logging
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import load_workbook
from playwright.async_api import BrowserContext, Locator, Page, async_playwright

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
LOG_DIR, SCREENSHOT_DIR = ROOT / "logs", ROOT / "screenshots"
SESSION_FILE = Path(__file__).with_name("session.json")

# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 默认按已确认存在数据的工位拆分导出，避免 Power BI 单文件 15 万行上限截断后部工位。
DEFAULT_EXPORT_STATIONS = ("15600", "20800", "23900", "24200", "24850", "60400")


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 支持通过环境变量调整工位批次，同时对空配置回退到当前已确认的工位清单。
def configured_export_stations() -> list[str]:
    configured = os.getenv("EXPORT_STATIONS", "").strip()
    stations = [item.strip() for item in configured.split(",") if item.strip()]
    return stations or list(DEFAULT_EXPORT_STATIONS)


# CHANGE [2026-09-10 09:15 +08:00] [WH400]: 内容校验失败属于确定性结果，不应重新登录并重复下载相同错误文件。
class NonRetryableExportError(RuntimeError):
    pass


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 将命中行数上限单独分类，使当前登录会话可自动改按日期二分，而不是终止整个工位任务。
class ExportRowLimitError(NonRetryableExportError):
    pass


# CHANGE [2026-09-09 15:18 +08:00] [WH400]: 日志同时写入终端与按日文件，且不记录账号密码等敏感配置。
def configure_logging() -> None:
    LOG_DIR.mkdir(exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.StreamHandler(), logging.FileHandler(LOG_DIR / f"export-{date.today():%Y%m%d}.log", encoding="utf-8")])


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 截图仅用于审计证据，限制等待时间并降级为警告，避免字体或渲染卡顿中断真实数据导出。
async def safe_screenshot(page: Page, path: Path) -> None:
    try:
        await page.screenshot(path=path, timeout=8000)
    except Exception as exc:
        logging.warning("证据截图失败但不影响导出：%s（%s）", path, exc)


# CHANGE [2026-09-09 15:18 +08:00] [WH400]: 优先按可访问文本定位 Power BI 控件，坐标后备留待真实报表首次校准。
async def click_text(page: Page, text: str, timeout: int = 60000) -> None:
    locator = page.get_by_text(text, exact=True).first
    await locator.wait_for(state="visible", timeout=timeout)
    await locator.click()


# CHANGE [2026-09-09 16:39 +08:00] [WH400]: Power BI 控件可能位于多层 iframe，循环所有 Frame 查找可见文本元素。
async def find_text_in_frames(page: Page, text: str, timeout: int = 60000) -> Locator:
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        for frame in page.frames:
            locator = frame.get_by_text(text, exact=True).first
            try:
                if await locator.is_visible():
                    return locator
            except Exception:
                continue
        await asyncio.sleep(0.5)
    raise RuntimeError(f"主页面及 iframe 中均未找到可见控件：{text}")


# CHANGE [2026-09-09 16:39 +08:00] [WH400]: 等待平台“正在加载数据”遮罩消失，避免固定等待不足时提前操作报表。
async def wait_for_report_ready(page: Page, timeout: int = 180000) -> None:
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        loading = False
        for frame in page.frames:
            indicators = frame.get_by_text("正在加载数据", exact=False)
            try:
                if await indicators.count() and await indicators.first.is_visible():
                    loading = True
                    break
            except Exception:
                continue
        if not loading:
            return
        await asyncio.sleep(1)
    raise RuntimeError("报表加载超过 180 秒，仍显示“正在加载数据”")


# CHANGE [2026-09-10 08:44 +08:00] [WH400]: 跨 iframe 收集日期输入框并选择最右侧的结束日期，避免误填开始日期或其他筛选项。
async def find_end_date_input(page: Page, timeout: int = 60000) -> Locator:
    selector = "input[type='date'], input[aria-label*='结束'], input[aria-label*='End'], input[placeholder*='结束']"
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        visible: list[tuple[float, Locator]] = []
        for frame in page.frames:
            inputs = frame.locator(selector)
            try:
                for index in range(await inputs.count()):
                    candidate = inputs.nth(index)
                    if await candidate.is_visible():
                        box = await candidate.bounding_box()
                        if box:
                            visible.append((box["x"] + box["width"] / 2, candidate))
            except Exception:
                continue
        if visible:
            return max(visible, key=lambda item: item[0])[1]
        await asyncio.sleep(0.5)
    raise RuntimeError("拧紧数据明细页中未找到 ResultTime 结束日期输入框")


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 同时定位 ResultTime 起止日期，供超限工位在保持完整日期覆盖的前提下自动拆分导出。
async def find_date_inputs(page: Page, timeout: int = 60000) -> list[Locator]:
    selector = "input[type='date'], input[aria-label*='日期'], input[aria-label*='Date']"
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        visible: list[tuple[float, Locator]] = []
        for frame in page.frames:
            inputs = frame.locator(selector)
            try:
                for index in range(await inputs.count()):
                    candidate = inputs.nth(index)
                    if await candidate.is_visible():
                        box = await candidate.bounding_box()
                        if box:
                            visible.append((box["x"] + box["width"] / 2, candidate))
            except Exception:
                continue
        if len(visible) >= 2:
            return [item[1] for item in sorted(visible, key=lambda item: item[0])[-2:]]
        await asyncio.sleep(0.5)
    raise RuntimeError("拧紧数据明细页中未找到完整的 ResultTime 起止日期输入框")


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 兼容斜杠及 ISO 日期回显，统一转换为日期对象用于无重叠的时间区间二分。
def parse_display_date(value: str) -> date:
    numbers = [int(item) for item in re.findall(r"\d+", value)]
    if len(numbers) < 3:
        raise RuntimeError(f"无法识别日期控件值：{value or '空'}")
    return date(numbers[0], numbers[1], numbers[2])


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 先扩展结束日期再提交开始日期，避免切换相邻区间时出现开始日期暂时晚于结束日期而被控件拒绝。
async def apply_date_range(report: Page, start: date, end: date) -> None:
    inputs = await find_date_inputs(report)
    current_end = await inputs[1].input_value()
    slash_format = "/" in current_end

    def display(value: date) -> str:
        return f"{value.year}/{value.month}/{value.day}" if slash_format else value.isoformat()

    await inputs[1].fill(display(end))
    await inputs[1].press("Tab")
    await report.wait_for_timeout(1800)
    await wait_for_report_ready(report)
    inputs = await find_date_inputs(report)
    await inputs[0].fill(display(start))
    await inputs[0].press("Tab")
    await report.wait_for_timeout(1800)
    await wait_for_report_ready(report)
    inputs = await find_date_inputs(report)
    actual_start = parse_display_date(await inputs[0].input_value())
    actual_end = parse_display_date(await inputs[1].input_value())
    if (actual_start, actual_end) != (start, end):
        raise RuntimeError(
            f"ResultTime 区间设置后未正确回显：期望 {start} 至 {end}，实际 {actual_start} 至 {actual_end}"
        )
    logging.info("ResultTime 区间已设置并校验：%s 至 %s", start, end)


# CHANGE [2026-09-10 08:44 +08:00] [WH400]: 将页签切换独立为前置步骤，确保日期筛选发生在“拧紧数据明细”页面而非默认趋势页。
async def open_detail_tab(report: Page) -> None:
    viewport = report.viewport_size or {"width": 1280, "height": 720}
    try:
        tab = await find_text_in_frames(report, "拧紧数据明细", timeout=15000)
        await tab.click()
    except RuntimeError:
        logging.info("未发现明细页签文本，按顶部第三页签截图坐标打开拧紧数据明细")
        await report.mouse.click(viewport["width"] * 0.59, viewport["height"] * 0.135)
    await report.wait_for_timeout(3000)
    await wait_for_report_ready(report)
    await safe_screenshot(report, SCREENSHOT_DIR / "05-detail-tab.png")


# CHANGE [2026-09-10 09:41 +08:00] [WH400]: 收集同名可见文本并选择页面最上方项，避免将表格列名误当作筛选器标题。
async def find_filter_label(page: Page, names: tuple[str, ...], timeout: int = 15000) -> tuple[Locator, str]:
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        for frame in page.frames:
            visible: list[tuple[float, int, Locator, str]] = []
            for name in names:
                locators = frame.get_by_text(name, exact=True)
                for index in range(await locators.count()):
                    locator = locators.nth(index)
                    try:
                        if await locator.is_visible():
                            box = await locator.bounding_box()
                            if box:
                                visible.append((box["y"], names.index(name), locator, name))
                    except Exception:
                        continue
            if visible:
                _, _, locator, name = min(visible, key=lambda item: (item[0], item[1]))
                return locator, name
        await asyncio.sleep(0.5)
    raise RuntimeError(f"未找到筛选器：{'、'.join(names)}")


# CHANGE [2026-09-10 09:41 +08:00] [WH400]: 返回筛选器容器与橡皮擦坐标，供点击后在同一容器内验证是否已显示“所有”。
async def slicer_clear_point(label: Locator) -> tuple[float, float, dict[str, float]]:
    label_box = await label.bounding_box()
    if not label_box:
        raise RuntimeError("筛选器标题没有可用的页面坐标")
    center_x = label_box["x"] + label_box["width"] / 2
    center_y = label_box["y"] + label_box["height"] / 2
    candidates: list[tuple[float, dict[str, float]]] = []
    ancestors = label.locator("xpath=ancestor::*")
    for index in range(await ancestors.count()):
        try:
            box = await ancestors.nth(index).bounding_box()
        except Exception:
            continue
        if not box:
            continue
        contains_label = (
            box["x"] <= center_x <= box["x"] + box["width"]
            and box["y"] <= center_y <= box["y"] + box["height"]
        )
        if contains_label and 110 <= box["width"] <= 500 and 40 <= box["height"] <= 180:
            candidates.append((box["width"] * box["height"], box))
    if candidates:
        box = min(candidates, key=lambda item: item[0])[1]
        return box["x"] + box["width"] - 22, center_y, box
    fallback_box = {
        "x": label_box["x"] - 10,
        "y": label_box["y"] - 12,
        "width": max(150, label_box["width"] + 90),
        "height": 70,
    }
    return fallback_box["x"] + fallback_box["width"] - 22, center_y, fallback_box


# CHANGE [2026-09-10 09:41 +08:00] [WH400]: 点击橡皮擦后要求同一筛选器容器显示“所有”，防止误点后继续导出带旧条件的数据。
async def clear_slicer_filter(page: Page, names: tuple[str, ...], required: bool = True) -> str | None:
    try:
        label, matched_name = await find_filter_label(page, names)
    except RuntimeError:
        if required:
            raise
        return None
    clear_x, clear_y, slicer_box = await slicer_clear_point(label)
    logging.info("清除筛选器 %s：悬浮橡皮擦位置 x=%.0f, y=%.0f", matched_name, clear_x, clear_y)
    await page.mouse.move(clear_x, clear_y)
    await page.wait_for_timeout(800)
    visible: list[tuple[float, Locator]] = []
    for frame in page.frames:
        buttons = frame.locator(
            "[aria-label*='清除'], [title*='清除'], [aria-label*='Clear'], [title*='Clear']"
        )
        for index in range(await buttons.count()):
            button = buttons.nth(index)
            try:
                if await button.is_visible():
                    box = await button.bounding_box()
                    if box:
                        distance = (
                            (box["x"] + box["width"] / 2 - clear_x) ** 2
                            + (box["y"] + box["height"] / 2 - clear_y) ** 2
                        )
                        visible.append((distance, button))
            except Exception:
                continue
    if visible:
        await min(visible, key=lambda item: item[0])[1].click()
    else:
        logging.info("筛选器橡皮擦未暴露可访问名称，按悬浮位置后备点击")
        await page.mouse.click(clear_x, clear_y)
    await page.wait_for_timeout(1200)
    await wait_for_report_ready(page)
    cleared = False
    for frame in page.frames:
        for all_text in ("所有", "All"):
            values = frame.get_by_text(all_text, exact=True)
            for index in range(await values.count()):
                try:
                    box = await values.nth(index).bounding_box()
                except Exception:
                    continue
                if not box:
                    continue
                value_x = box["x"] + box["width"] / 2
                value_y = box["y"] + box["height"] / 2
                if (
                    slicer_box["x"] <= value_x <= slicer_box["x"] + slicer_box["width"]
                    and slicer_box["y"] <= value_y <= slicer_box["y"] + slicer_box["height"]
                ):
                    cleared = True
                    break
            if cleared:
                break
        if cleared:
            break
    if not cleared:
        raise RuntimeError(f"筛选器 {matched_name} 点击橡皮擦后未显示“所有”，已停止后续导出")
    logging.info("筛选器 %s 已清除并校验为“所有”", matched_name)
    return matched_name


# CHANGE [2026-09-10 09:41 +08:00] [WH400]: 优先识别合并筛选器并只清除一次；分开展示时再分别清除电枪工位和程序名称。
async def clear_detail_filters(report: Page) -> None:
    station_name = await clear_slicer_filter(report, ("电枪工位、程序名称", "电枪工位"))
    if station_name != "电枪工位、程序名称":
        await clear_slicer_filter(report, ("程序名称",))
    await clear_slicer_filter(report, ("结果状态",))
    await safe_screenshot(report, SCREENSHOT_DIR / "06-filters-cleared.png")


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 在工位筛选器下拉层内按角色或邻近位置选择目标工位，并校验筛选器回显，防止误点明细表中的同名值。
async def select_station_filter(page: Page, station: str) -> None:
    label, matched_name = await find_filter_label(page, ("电枪工位、程序名称", "电枪工位"))
    _, _, slicer_box = await slicer_clear_point(label)
    dropdown_x = slicer_box["x"] + slicer_box["width"] / 2
    dropdown_y = slicer_box["y"] + slicer_box["height"] - 18
    logging.info("选择工位 %s：打开筛选器 %s", station, matched_name)
    await page.mouse.click(dropdown_x, dropdown_y)
    await page.wait_for_timeout(700)

    # CHANGE [2026-09-11 10:23 +08:00] [WH400]: Power BI 下拉列表采用虚拟滚动，优先填写同列搜索框以加载不在首屏的工位。
    search_candidates: list[tuple[float, Locator]] = []
    for frame in page.frames:
        searches = frame.locator(
            "input[placeholder*='搜索'], input[aria-label*='搜索'], "
            "input[placeholder*='Search'], input[aria-label*='Search']"
        )
        try:
            for index in range(await searches.count()):
                search = searches.nth(index)
                if not await search.is_visible():
                    continue
                box = await search.bounding_box()
                if not box:
                    continue
                center_x = box["x"] + box["width"] / 2
                if slicer_box["x"] - 30 <= center_x <= slicer_box["x"] + slicer_box["width"] + 30:
                    search_candidates.append((abs(box["y"] - (slicer_box["y"] + slicer_box["height"])), search))
        except Exception:
            continue
    if search_candidates:
        search = min(search_candidates, key=lambda item: item[0])[1]
        await search.fill(station)
        await page.wait_for_timeout(900)

    candidates: list[tuple[int, float, Locator]] = []
    for frame in page.frames:
        tree_items = frame.locator(f'[role="treeitem"][title="{station}"]')
        try:
            for index in range(await tree_items.count()):
                item = tree_items.nth(index)
                if await item.is_visible():
                    box = await item.bounding_box()
                    if box:
                        candidates.append((-1, abs(box["y"] - (slicer_box["y"] + slicer_box["height"])), item))
        except Exception:
            pass
        role_options = frame.get_by_role("option", name=station, exact=True)
        try:
            for index in range(await role_options.count()):
                option = role_options.nth(index)
                if await option.is_visible():
                    box = await option.bounding_box()
                    if box:
                        candidates.append((0, abs(box["y"] - (slicer_box["y"] + slicer_box["height"])), option))
        except Exception:
            continue
    # CHANGE [2026-09-11 10:23 +08:00] [WH400]: 仅把下拉 treeitem 作为文本后备，排除明细表同名值，并通过逐屏滚动加载非首屏工位。
    if not candidates:
        viewport = page.viewport_size or {"width": 1280, "height": 720}
        scroll_y = min(viewport["height"] - 50, slicer_box["y"] + slicer_box["height"] + 100)
        for _ in range(12):
            await page.mouse.move(dropdown_x, scroll_y)
            await page.mouse.wheel(0, 420)
            await page.wait_for_timeout(500)
            for frame in page.frames:
                tree_items = frame.locator(f'[role="treeitem"][title="{station}"]')
                try:
                    for index in range(await tree_items.count()):
                        item = tree_items.nth(index)
                        if await item.is_visible():
                            box = await item.bounding_box()
                            if box:
                                candidates.append((-1, abs(box["y"] - scroll_y), item))
                except Exception:
                    pass
                role_options = frame.get_by_role("option", name=station, exact=True)
                try:
                    for index in range(await role_options.count()):
                        option = role_options.nth(index)
                        if await option.is_visible():
                            box = await option.bounding_box()
                            if box:
                                candidates.append((0, abs(box["y"] - scroll_y), option))
                except Exception:
                    continue
                exact_values = frame.get_by_text(station, exact=True)
                try:
                    for index in range(await exact_values.count()):
                        value = exact_values.nth(index)
                        if not await value.is_visible():
                            continue
                        tree_parent = value.locator("xpath=ancestor::*[@role='treeitem'][1]")
                        if await tree_parent.count() and await tree_parent.is_visible():
                            box = await tree_parent.bounding_box()
                            if box:
                                candidates.append((-1, abs(box["y"] - scroll_y), tree_parent))
                except Exception:
                    continue
            if candidates:
                break
    if not candidates:
        await safe_screenshot(page, SCREENSHOT_DIR / f"07-station-{station}-not-found.png")
        await page.keyboard.press("Escape")
        raise NonRetryableExportError(f"工位筛选器下拉列表中未找到工位 {station}")
    await min(candidates, key=lambda item: (item[0], item[1]))[2].click()
    await page.wait_for_timeout(1800)
    await wait_for_report_ready(page)

    selected = False
    for frame in page.frames:
        values = frame.get_by_text(station, exact=True)
        for index in range(await values.count()):
            try:
                box = await values.nth(index).bounding_box()
            except Exception:
                continue
            if not box:
                continue
            center_x = box["x"] + box["width"] / 2
            center_y = box["y"] + box["height"] / 2
            if (
                slicer_box["x"] <= center_x <= slicer_box["x"] + slicer_box["width"]
                and slicer_box["y"] <= center_y <= slicer_box["y"] + slicer_box["height"]
            ):
                selected = True
                break
        if selected:
            break
    if not selected:
        raise NonRetryableExportError(f"工位 {station} 选择后未在筛选器中正确回显")
    logging.info("工位筛选已设置并校验：%s", station)
    await safe_screenshot(page, SCREENSHOT_DIR / f"07-station-{station}-selected.png")


# CHANGE [2026-09-10 09:41 +08:00] [WH400]: 仅接受明细表目标点 90 像素内的省略号，忽略清除筛选后仍可见的筛选器菜单。
async def find_more_options(page: Page, target_x: float, target_y: float, timeout: int = 15000) -> Locator:
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        visible: list[tuple[float, Locator]] = []
        for frame in page.frames:
            candidates = frame.locator(
                "[aria-label*='更多选项'], [title*='更多选项'], "
                "[aria-label*='More options'], [title*='More options']"
            )
            for index in range(await candidates.count()):
                candidate = candidates.nth(index)
                try:
                    if await candidate.is_visible():
                        box = await candidate.bounding_box()
                        if box:
                            center_x = box["x"] + box["width"] / 2
                            center_y = box["y"] + box["height"] / 2
                            distance = (center_x - target_x) ** 2 + (center_y - target_y) ** 2
                            visible.append((distance, candidate))
                except Exception:
                    continue
        if visible:
            distance, candidate = min(visible, key=lambda item: item[0])
            if distance ** 0.5 <= 90:
                return candidate
        await asyncio.sleep(0.5)
    raise RuntimeError("明细表附近未找到“更多选项”按钮")


# CHANGE [2026-09-10 09:52 +08:00] [WH400]: 跨 iframe 检测菜单文本及搜索框，用于区分明细表菜单与筛选器菜单。
async def inspect_export_menu(page: Page) -> tuple[bool, str]:
    visible_texts: set[str] = set()
    for frame in page.frames:
        for text in ("导出数据", "以表的形式显示", "显示为表", "聚焦", "聚焦模式", "排列轴"):
            locator = frame.get_by_text(text, exact=True)
            try:
                if await locator.count() and await locator.first.is_visible():
                    visible_texts.add(text)
            except Exception:
                continue
    search_visible = False
    for frame in page.frames:
        search = frame.locator("input[placeholder*='搜索'], input[aria-label*='搜索'], input[placeholder*='Search']")
        try:
            if await search.count() and await search.first.is_visible():
                search_visible = True
                break
        except Exception:
            continue
    wrong_markers = []
    if "排列轴" in visible_texts:
        wrong_markers.append("排列轴")
    if search_visible:
        wrong_markers.append("搜索框")
    has_export = "导出数据" in visible_texts
    if has_export and not wrong_markers:
        return True, "菜单包含导出数据，且未发现筛选器专属控件"
    detail = "、".join(wrong_markers) if wrong_markers else "未发现导出数据"
    return False, detail


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 为每个工位保留独立工具栏与菜单证据截图，便于定位某一批次的视觉对象误点问题。
async def open_detail_export_menu(page: Page, evidence_prefix: str = "detail") -> None:
    viewport = page.viewport_size or {"width": 1280, "height": 720}
    expected_x, expected_y = viewport["width"] * 0.805, viewport["height"] * 0.34
    hover_ratios = (
        (0.805, 0.34),
        (0.79, 0.36),
        (0.82, 0.36),
        (0.75, 0.45),
    )
    for attempt, (x_ratio, y_ratio) in enumerate(hover_ratios, 1):
        await page.keyboard.press("Escape")
        await page.mouse.move(viewport["width"] * 0.55, viewport["height"] * 0.60)
        await page.wait_for_timeout(400)
        target_x, target_y = viewport["width"] * x_ratio, viewport["height"] * y_ratio
        logging.info("第 %s 个位置悬浮明细表工具栏：x=%.0f, y=%.0f", attempt, target_x, target_y)
        await page.mouse.move(target_x, target_y)
        await page.wait_for_timeout(1200)
        try:
            more = await find_more_options(page, expected_x, expected_y, timeout=2500)
        except RuntimeError:
            continue
        # CHANGE [2026-09-28 15:42 +08:00] [WH400]: 找到动态工具栏按钮后立即直接点击元素，避免截图触发 Power BI 重绘后旧坐标落空。
        try:
            await more.click(timeout=2500, force=True)
        except Exception as exc:
            logging.warning("第 %s 次点击明细表更多选项失败：%s", attempt, exc)
            continue
        await page.wait_for_timeout(900)
        await safe_screenshot(page, SCREENSHOT_DIR / f"09-{evidence_prefix}-toolbar.png")
        valid, detail = await inspect_export_menu(page)
        if valid:
            logging.info("明细表菜单校验通过：%s", detail)
            await safe_screenshot(page, SCREENSHOT_DIR / f"10-{evidence_prefix}-export-menu.png")
            return
        logging.warning("第 %s 次菜单校验失败（%s），关闭菜单后继续尝试", attempt, detail)
        await safe_screenshot(page, SCREENSHOT_DIR / f"10-{evidence_prefix}-wrong-menu-{attempt}.png")
        await page.keyboard.press("Escape")
        await page.wait_for_timeout(500)
    raise NonRetryableExportError("未能打开通过内容校验的明细表导出菜单；已停止且不重新登录重试")


# CHANGE [2026-09-10 09:15 +08:00] [WH400]: 重置 Power BI 错写的 A1 工作表尺寸并兼容“字段 的总和”表头，防止 5 万行正常文件被误判。
def validate_export_workbook(path: Path) -> tuple[bool, str]:
    required = {
        "电枪工位", "程序名称", "ESN", "扭矩上限", "扭矩下限", "最终扭矩",
        "角度上限", "角度下限", "最终角度", "结果状态", "ResultTime", "螺栓",
    }

    def normalize(value: object) -> str:
        text = re.sub(r"\s+", "", "" if value is None else str(value)).lstrip("\ufeff")
        return re.sub(r"的?(总和|平均值|计数|最大值|最小值)$", "", text)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            if sheet.calculate_dimension() == "A1:A1":
                sheet.reset_dimensions()
            preview: list[str] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_index > 20:
                    break
                headers = {normalize(value) for value in row if value is not None}
                preview.extend(headers)
                if required.issubset(headers):
                    return True, f"{sheet.title} 已识别完整明细表头"
        if any("应用的筛选器" in value for value in preview):
            return False, "导出内容是 Power BI“应用的筛选器”摘要，不是拧紧数据明细表"
        return False, "导出文件前 20 行未识别到完整的 12 列明细表头"
    finally:
        workbook.close()


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 逐批核对导出文件只含目标工位且未逼近 Excel 导出上限，避免错误筛选或静默截断进入解析平台。
def validate_export_station(path: Path, expected_station: str) -> tuple[bool, str, int]:
    def normalize(value: object) -> str:
        text = re.sub(r"\s+", "", "" if value is None else str(value)).lstrip("\ufeff")
        return re.sub(r"的?(总和|平均值|计数|最大值|最小值)$", "", text)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            if sheet.calculate_dimension() == "A1:A1":
                sheet.reset_dimensions()
            station_index: int | None = None
            data_rows = 0
            unexpected: set[str] = set()
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if station_index is None:
                    headers = [normalize(value) for value in row]
                    if "电枪工位" in headers:
                        station_index = headers.index("电枪工位")
                    elif row_index >= 20:
                        break
                    continue
                value = row[station_index] if station_index < len(row) else None
                station_value = normalize(value)
                if not station_value:
                    continue
                data_rows += 1
                if station_value != expected_station:
                    unexpected.add(station_value)
                    if len(unexpected) >= 5:
                        break
            if station_index is None:
                continue
            if unexpected:
                return False, f"文件含非目标工位：{'、'.join(sorted(unexpected))}", data_rows
            if data_rows == 0:
                return False, f"文件中没有工位 {expected_station} 的数据", 0
            if data_rows >= 149900:
                return False, f"工位 {expected_station} 导出 {data_rows} 行，仍可能触发 15 万行截断", data_rows
            return True, f"仅包含工位 {expected_station}，共 {data_rows} 条", data_rows
        return False, "未识别到电枪工位列", 0
    finally:
        workbook.close()


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 全自动模式兼容已有 OA 会话，并在同窗多用户提示出现时先注销冲突会话再重新登录。
async def ensure_login(page: Page, mode: str) -> None:
    await page.goto(os.getenv("OA_URL", "https://oa.xcec.com.cn:8080/seeyon/index.jsp"), wait_until="domcontentloaded")
    if mode == "auto":
        username, password = os.getenv("OA_USERNAME"), os.getenv("OA_PASSWORD")
        if not username or not password:
            raise RuntimeError("全自动模式需要在 .env 中配置 OA_USERNAME 与 OA_PASSWORD")
        platform_entry = page.get_by_text("智能数据平台", exact=True)
        deadline = asyncio.get_running_loop().time() + 60
        while asyncio.get_running_loop().time() < deadline:
            conflict_logged_out = False
            for frame in page.frames:
                logout = frame.get_by_text("注销", exact=True)
                try:
                    if await logout.count() and await logout.first.is_visible():
                        logging.info("检测到 OA 同一窗口多用户提示，点击注销后重新登录")
                        await logout.first.click()
                        conflict_logged_out = True
                        break
                except Exception:
                    continue
            if conflict_logged_out:
                await page.wait_for_timeout(1800)
                continue
            try:
                if await platform_entry.count() and await platform_entry.first.is_visible():
                    logging.info("检测到有效 OA 会话，跳过账号密码填写")
                    break
            except Exception:
                pass
            # CHANGE [2026-09-11 10:23 +08:00] [WH400]: OA 版本可能把登录表单放入 iframe，逐帧寻找成对的账号和密码输入框。
            login_submitted = False
            for frame in page.frames:
                username_input = frame.locator("input[type='text']").first
                password_input = frame.locator("input[type='password']").first
                try:
                    if await username_input.is_visible() and await password_input.is_visible():
                        await username_input.fill(username)
                        await password_input.fill(password)
                        login_button = frame.get_by_role("button", name="登录")
                        if await login_button.count():
                            await login_button.first.click()
                        else:
                            await password_input.press("Enter")
                        login_submitted = True
                        break
                except Exception:
                    continue
            if login_submitted:
                await page.wait_for_timeout(1800)
                continue
            await asyncio.sleep(0.5)
        else:
            logging.info("未发现登录输入框，继续等待已有 OA 会话完成首页加载")
    else:
        logging.info("半自动模式：请在浏览器中完成登录，脚本将等待智能数据平台入口出现")
    entry_state = "attached" if mode == "auto" else "visible"
    await page.get_by_text("智能数据平台", exact=True).wait_for(state=entry_state, timeout=300000)


# CHANGE [2026-09-09 16:32 +08:00] [WH400]: 真实环境域名与需求文档拼写不同，综合主机名和固定路径识别数据平台。
def is_platform_url(url: str) -> bool:
    parsed = urlparse(url)
    configured_host = urlparse(os.getenv("DATA_PLATFORM_URL", "")).hostname or ""
    hostname = (parsed.hostname or "").lower()
    return (
        hostname == configured_host.lower()
        or hostname.startswith("databi.")
        or ("xcec.com.cn" in hostname and parsed.path.rstrip("/").endswith("/page/home"))
    )


# CHANGE [2026-09-09 16:32 +08:00] [WH400]: 按真实页面特征等待数据平台，兼容 OA 在当前页或新标签页打开入口。
async def wait_for_platform_page(context: BrowserContext, timeout: int = 60000) -> Page:
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        for candidate in reversed(context.pages):
            if is_platform_url(candidate.url):
                await candidate.wait_for_load_state("domcontentloaded")
                return candidate
        await asyncio.sleep(0.5)
    urls = "、".join(candidate.url for candidate in context.pages)
    raise RuntimeError(f"点击智能数据平台后未识别到平台页面；当前页面：{urls}")


# CHANGE [2026-09-09 16:24 +08:00] [WH400]: 报表入口同时兼容弹出新标签、当前页导航和 SPA 内嵌加载三种行为。
async def wait_for_report_page(context: BrowserContext, platform: Page, previous_url: str,
                               previous_pages: set[Page], timeout: int = 60000) -> Page:
    deadline = asyncio.get_running_loop().time() + timeout / 1000
    while asyncio.get_running_loop().time() < deadline:
        new_pages = [candidate for candidate in context.pages if candidate not in previous_pages]
        if new_pages:
            report = new_pages[-1]
            await report.wait_for_load_state("domcontentloaded")
            return report
        if platform.url != previous_url:
            await platform.wait_for_load_state("domcontentloaded")
            return platform
        for marker in ("筛选器", "拧紧数据明细", "拧紧趋势监控"):
            if await platform.get_by_text(marker, exact=True).count():
                return platform
        await asyncio.sleep(0.5)
    raise RuntimeError("点击拧紧数据报表后，未检测到报表页面或页签")


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: OA 自动登录后入口可能折叠为隐藏 DOM，短暂等待不可见时直接触发已挂载入口的点击事件。
async def open_report(context: BrowserContext, page: Page) -> Page:
    logging.info("步骤 2/6：点击智能数据平台")
    try:
        await click_text(page, "智能数据平台", timeout=8000)
    except Exception:
        entry = page.get_by_text("智能数据平台", exact=True).first
        await entry.wait_for(state="attached", timeout=30000)
        await entry.evaluate("element => element.click()")
    platform = await wait_for_platform_page(context)
    logging.info("已进入数据平台：%s", platform.url)
    await safe_screenshot(platform, SCREENSHOT_DIR / "02-platform-home.png")

    logging.info("步骤 3/6：展开质量部并进入工厂质量")
    quality = platform.get_by_text("质量部", exact=True).first
    await quality.wait_for(state="visible", timeout=60000)
    await quality.hover()
    factory = platform.get_by_text("工厂质量", exact=True).first
    try:
        await factory.wait_for(state="visible", timeout=5000)
    except Exception:
        await quality.click()
        await factory.wait_for(state="visible", timeout=15000)
    await factory.click()

    logging.info("步骤 4/6：等待并打开拧紧数据-近半年")
    card = platform.get_by_text("拧紧数据-近半年", exact=True).first
    await card.wait_for(state="visible", timeout=60000)
    await safe_screenshot(platform, SCREENSHOT_DIR / "03-factory-quality.png")
    previous_url = platform.url
    previous_pages = set(context.pages)
    await card.click()
    report = await wait_for_report_page(context, platform, previous_url, previous_pages)
    await report.wait_for_timeout(8000)
    return report


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 设置昨天前先保存报表当前开始日期，供超限工位自动覆盖完整原始时间范围。
async def apply_yesterday_filter(report: Page) -> tuple[date, date]:
    yesterday = date.today() - timedelta(days=1)
    await wait_for_report_ready(report)
    await find_text_in_frames(report, "ResultTime", timeout=60000)
    date_inputs = await find_date_inputs(report)
    original_start = parse_display_date(await date_inputs[0].input_value())
    end_date = await find_end_date_input(report)
    previous_value = await end_date.input_value()
    display_value = (
        f"{yesterday.year}/{yesterday.month}/{yesterday.day}"
        if "/" in previous_value
        else yesterday.isoformat()
    )
    logging.info("ResultTime 结束日期提交：%s → %s", previous_value, display_value)
    await end_date.fill(display_value)
    await end_date.press("Tab")
    await report.wait_for_timeout(2500)
    await wait_for_report_ready(report)
    refreshed = await find_end_date_input(report)
    value = await refreshed.input_value()
    numbers = [int(item) for item in re.findall(r"\d+", value)]
    if len(numbers) < 3 or tuple(numbers[:3]) != (yesterday.year, yesterday.month, yesterday.day):
        logging.info("日期首次填充未生效，改用全选逐字输入后以 Tab 提交")
        await refreshed.click()
        await refreshed.press("Control+A")
        await refreshed.type(display_value, delay=80)
        await refreshed.press("Tab")
        await report.wait_for_timeout(2500)
        await wait_for_report_ready(report)
        refreshed = await find_end_date_input(report)
        value = await refreshed.input_value()
        numbers = [int(item) for item in re.findall(r"\d+", value)]
    if len(numbers) < 3 or tuple(numbers[:3]) != (yesterday.year, yesterday.month, yesterday.day):
        raise RuntimeError(f"ResultTime 结束日期设置后未正确回显：期望 {yesterday.isoformat()}，实际 {value or '空'}")
    logging.info("拧紧数据明细页 ResultTime 结束日期已设置并校验：%s", value)
    await safe_screenshot(report, SCREENSHOT_DIR / "07-detail-filtered.png")
    return original_start, yesterday


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 导出文件名、截图和内容校验均绑定工位批次；日期二分出现无数据区间时保留证据并继续其余区间。
async def export_tab(report: Page, tab_name: str, download_dir: Path, station: str,
                     period: tuple[date, date] | None = None) -> Path | None:
    if tab_name != "拧紧数据明细":
        tab = await find_text_in_frames(report, tab_name, timeout=10000)
        await tab.click()
        await report.wait_for_timeout(3000)
        await wait_for_report_ready(report)
    evidence_prefix = f"station-{station}"
    await safe_screenshot(report, SCREENSHOT_DIR / f"08-{evidence_prefix}-detail-ready.png")

    await open_detail_export_menu(report, evidence_prefix)
    export_data = await find_text_in_frames(report, "导出数据", timeout=30000)
    await export_data.click()
    await report.wait_for_timeout(900)
    await safe_screenshot(report, SCREENSHOT_DIR / f"11-{evidence_prefix}-export-dialog.png")
    try:
        summary = await find_text_in_frames(report, "汇总数据", timeout=10000)
        await summary.click()
    except RuntimeError:
        logging.info("未发现汇总数据文本，保留弹窗默认选项")
    async with report.expect_download(timeout=120000) as download_info:
        export_button = await find_text_in_frames(report, "导出", timeout=30000)
        await export_button.click()
    download = await download_info.value
    download_dir.mkdir(parents=True, exist_ok=True)
    period_suffix = f"_{period[0]:%Y-%m-%d}_至_{period[1]:%Y-%m-%d}" if period else ""
    target = download_dir / f"{tab_name}_工位{station}{period_suffix}_汇总_{date.today() - timedelta(days=1):%Y-%m-%d}.xlsx"
    await download.save_as(target)
    if not target.exists() or target.stat().st_size == 0:
        raise RuntimeError(f"下载文件无效：{target}")
    valid, detail = validate_export_workbook(target)
    if not valid:
        rejected = target.with_name(f"错误导出_{datetime.now():%Y%m%d-%H%M%S}_{target.name}")
        target.replace(rejected)
        raise NonRetryableExportError(f"{detail}；错误文件已保留为：{rejected}")
    station_valid, station_detail, row_count = validate_export_station(target, station)
    if not station_valid:
        if row_count == 0 and "没有工位" in station_detail:
            empty = target.with_name(f"空数据_{target.name}")
            target.replace(empty)
            logging.info("工位 %s 在当前日期区间无数据，已跳过并保留证据：%s", station, empty)
            return None
        rejected = target.with_name(f"错误导出_{datetime.now():%Y%m%d-%H%M%S}_{target.name}")
        target.replace(rejected)
        if row_count >= 149900:
            raise ExportRowLimitError(f"{station_detail}；超限文件已保留为：{rejected}")
        raise NonRetryableExportError(f"{station_detail}；错误文件已保留为：{rejected}")
    logging.info("导出文件校验通过：%s（%s；%s；%s 条）", target, detail, station_detail, row_count)
    return target


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 工位全区间超限时递归二分日期且边界不重叠，直至每个文件均低于 Power BI 行数上限。
async def export_station_with_split(report: Page, tab_name: str, download_dir: Path, station: str,
                                    start: date, end: date, period: tuple[date, date] | None = None) -> list[Path]:
    try:
        output = await export_tab(report, tab_name, download_dir, station, period)
        return [output] if output else []
    except ExportRowLimitError as exc:
        split_start, split_end = period or (start, end)
        if split_start >= split_end:
            raise NonRetryableExportError(f"工位 {station} 单日数据仍超过导出上限：{split_start}") from exc
        midpoint = split_start + timedelta(days=(split_end - split_start).days // 2)
        left = (split_start, midpoint)
        right = (midpoint + timedelta(days=1), split_end)
        logging.warning("工位 %s 命中导出上限，自动拆分为 %s 至 %s、%s 至 %s", station, *left, *right)
        outputs: list[Path] = []
        for child_period in (left, right):
            await apply_date_range(report, *child_period)
            outputs.extend(await export_station_with_split(
                report, tab_name, download_dir, station, start, end, child_period
            ))
        return outputs


# CHANGE [2026-09-11 10:23 +08:00] [WH400]: 日期固定为昨天后按工位逐批清除、选择、导出和校验，绕开 Power BI 全量导出的行数上限。
async def run_once(mode: str) -> list[Path]:
    SCREENSHOT_DIR.mkdir(exist_ok=True)
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=os.getenv("HEADLESS", "false").lower() == "true")
        context = await browser.new_context(ignore_https_errors=True, storage_state=str(SESSION_FILE) if SESSION_FILE.exists() else None,
                                            accept_downloads=True)
        page = await context.new_page()
        await ensure_login(page, mode)
        await safe_screenshot(page, SCREENSHOT_DIR / "01-login-ok.png")
        await context.storage_state(path=SESSION_FILE)
        report = await open_report(context, page)
        await safe_screenshot(report, SCREENSHOT_DIR / "04-report-loaded.png")
        logging.info("步骤 5/6：进入拧紧数据明细并设置 ResultTime 结束日期为昨天")
        try:
            await open_detail_tab(report)
            logging.info("清除电枪工位、程序名称和结果状态筛选")
            await clear_detail_filters(report)
            original_start, yesterday = await apply_yesterday_filter(report)
        except Exception:
            await safe_screenshot(report, SCREENSHOT_DIR / "07-detail-filter-failed.png")
            raise
        tabs = [item.strip() for item in os.getenv("EXPORT_TABS", "拧紧数据明细").split(",") if item.strip()]
        stations = configured_export_stations()
        logging.info("步骤 6/6：按工位分批导出页签 %s；工位：%s", "、".join(tabs), "、".join(stations))
        try:
            outputs: list[Path] = []
            for station_index, station in enumerate(stations, 1):
                logging.info("开始工位批次 %s/%s：%s", station_index, len(stations), station)
                if station_index > 1:
                    await clear_slicer_filter(report, ("电枪工位、程序名称", "电枪工位"))
                await select_station_filter(report, station)
                for tab in tabs:
                    station_outputs = await export_station_with_split(
                        report, tab, ROOT / os.getenv("DOWNLOAD_DIR", "downloads"),
                        station, original_start, yesterday
                    )
                    outputs.extend(station_outputs)
                if station_index < len(stations):
                    await apply_date_range(report, original_start, yesterday)
        except Exception:
            await safe_screenshot(report, SCREENSHOT_DIR / "12-export-failed.png")
            raise
        await context.storage_state(path=SESSION_FILE)
        await browser.close()
        return outputs


# CHANGE [2026-09-10 09:15 +08:00] [WH400]: 浏览器暂态异常保留重试，但内容校验失败立即停止，避免重复登录和下载相同文件。
async def run_with_retry(mode: str, retries: int = 2) -> None:
    for attempt in range(1, retries + 2):
        try:
            outputs = await run_once(mode)
            logging.info("导出完成：%s", ", ".join(map(str, outputs)))
            return
        except NonRetryableExportError as exc:
            logging.error("导出文件内容校验未通过，不再重试：%s", exc)
            raise
        except Exception as exc:
            logging.exception("第 %s 次执行失败：%s", attempt, exc)
            if attempt > retries:
                raise


# CHANGE [2026-09-09 15:18 +08:00] [WH400]: 命令行入口支持用户显式选择全自动或半自动接管模式。
def main() -> None:
    parser = argparse.ArgumentParser(description="拧紧数据自动导出工具")
    parser.add_argument("--mode", choices=["auto", "semi-auto"], default="semi-auto")
    args = parser.parse_args()
    configure_logging()
    asyncio.run(run_with_retry(args.mode))


if __name__ == "__main__":
    main()
